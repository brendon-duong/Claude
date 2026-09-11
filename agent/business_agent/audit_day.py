"""One day's shift audit from the Zoom Phone call logs, in Elaine's layout.

    python3 -m business_agent.audit_day                 # today, Manila
    python3 -m business_agent.audit_day 2026-09-10 --xlsx out.xlsx --json out.json

This is the half of the audit the phone system can answer: how many calls,
how many answered, how many long enough to be a survey, and how much of the
three hours was spent away from the phone. It cannot answer what the caller
*declared* — that comes from #results-pacificlinkglobal on Slack, and the JSON
this writes is the shape a later step merges those numbers into.

Elaine's sheet has seven columns. Three of them come from the caller's own
report and are left blank here rather than guessed: completed surveys as
declared, total calls as declared, and Discrepancy Y/N, which is a comparison
of the two halves and not a call-log judgement.

Everything here is testable without Zoom: `audit()` takes the same `fetch`
stub `calls_for_day` does.
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta

from .calllogs import (
    COMPLETE_SECONDS,
    IDLE_SECONDS,
    MANILA,
    SHIFT_HOURS,
    ShiftCalls,
    by_caller,
    calls_for_day,
)
from .names import Match, match

# Elaine lists a break in the Details column when it is this long or longer.
# Shorter gaps still count toward cumulative off-phone time; they are just not
# worth a line each.
LISTED_BREAK_SECONDS = 300

# A completed survey within this much of the threshold is worth a mention —
# Elaine flags "1 completed survey was below 2 min and 30 seconds", and a call
# that only just cleared it is the same kind of borderline.
NEAR_THRESHOLD_SECONDS = 30


@dataclass(frozen=True)
class Row:
    """One caller's day, as the phone system saw it."""

    who: Match
    shift: ShiftCalls

    @property
    def name(self) -> str:
        """The roster name if settled, otherwise the Zoom name, marked."""
        return self.who.roster or f"{self.who.zoom} (?)"

    @property
    def details(self) -> list[str]:
        """Elaine's Details column, from the call-log side only."""
        lines = []
        for start, length in sorted(self.shift.breaks(), key=lambda b: -b[1].total_seconds()):
            if length.total_seconds() >= LISTED_BREAK_SECONDS:
                mins = int(length.total_seconds() // 60)
                lines.append(
                    f"{mins} min break between {_clock(start)} - {_clock(start + length)}"
                )
        idle = int(self.shift.idle_time().total_seconds() // 60)
        if idle:
            lines.append(f"{idle} min cumulative off-phone time throughout shift")
        near = sum(
            1 for c in self.shift.calls
            if c.answered and COMPLETE_SECONDS <= c.seconds < COMPLETE_SECONDS + NEAR_THRESHOLD_SECONDS
        )
        if near:
            plural = "s" if near > 1 else ""
            lines.append(
                f"{near} completed survey{plural} within {NEAR_THRESHOLD_SECONDS} seconds "
                f"of the {COMPLETE_SECONDS // 60} min {COMPLETE_SECONDS % 60} threshold"
            )
        return lines


@dataclass(frozen=True)
class Report:
    day: date
    rows: tuple[Row, ...]          # worked the shift, sorted by completes
    off_shift: tuple[Row, ...]     # made calls, none inside the window — listed, not scored
    # Extensions that only received calls nobody answered. Not callers, not
    # listed by name: a missed call to a desk is not something the person did.
    rang_only: int = 0

    @property
    def needing_a_person(self) -> list[Row]:
        return [r for r in self.rows + self.off_shift if r.who.needs_a_person]


def _clock(when: datetime) -> str:
    """Manila, the way Elaine writes it: 2:18, not 14:18."""
    local = when.astimezone(MANILA)
    return f"{local.hour % 12 or 12}:{local.minute:02d}"


def _hm(td: timedelta) -> str:
    mins = int(td.total_seconds() // 60)
    return f"{mins // 60}h{mins % 60:02d}m" if mins >= 60 else f"{mins}m"


def today_manila() -> date:
    return datetime.now(MANILA).date()


def audit(day: date, *, fetch=None, token: str = "") -> Report:
    grouped = by_caller(calls_for_day(day, fetch=fetch, token=token), day)
    rows, off, rang_only = [], [], 0
    for shift in grouped.values():
        if not shift.activity:
            rang_only += 1
            continue
        row = Row(who=match(shift.caller), shift=shift)
        (rows if shift.on_shift else off).append(row)
    rows.sort(key=lambda r: (-r.shift.completes(), -r.shift.answered))
    off.sort(key=lambda r: r.name)
    return Report(day=day, rows=tuple(rows), off_shift=tuple(off), rang_only=rang_only)


def render_text(report: Report) -> str:
    out = []
    p = out.append
    p(f"{report.day.strftime('%A %-d %B %Y')} — from the Zoom Phone call logs")
    p(f"complete = answered and >= {COMPLETE_SECONDS}s · idle floor {IDLE_SECONDS}s · "
      f"shift {SHIFT_HOURS}h from own first call, floored 1:30pm Manila")
    p("")
    p(f"{'':1}{'Caller':<28}{'Calls':>6}{'Ans':>5}{'Comp':>6}{'Away':>7}{'Worked':>8}{'Short':>7}  Window")
    p("-" * 92)
    for r in report.rows:
        s = r.shift
        flag = "!" if s.shortfall() > timedelta(minutes=30) else " "
        p(f"{flag}{r.name[:27]:<28}{s.attempts:>6}{s.answered:>5}{s.completes():>6}"
          f"{_hm(s.idle_time()):>7}{_hm(s.worked()):>8}{_hm(s.shortfall()):>7}  "
          f"{_clock(s.started_at())}-{_clock(s.finished_at())}")
    p("-" * 92)
    p(f"{'':1}{'TOTAL':<28}{sum(r.shift.attempts for r in report.rows):>6}"
      f"{sum(r.shift.answered for r in report.rows):>5}{sum(r.shift.completes() for r in report.rows):>6}")
    over = [r for r in report.rows if r.shift.shortfall() > timedelta(minutes=30)]
    p(f"! = more than 30 minutes short — {len(over)} of {len(report.rows)}")
    if report.off_shift:
        p("")
        p("Made calls but never inside the shift window — listed, not scored:")
        for r in report.off_shift:
            last = max(c.started for c in r.shift.calls if c.presence)
            p(f"  {r.name:<28}{r.shift.activity} call(s), last at {_clock(last)} Manila")
    if report.rang_only:
        p("")
        p(f"{report.rang_only} extension(s) only received calls nobody answered — not listed.")
    need = report.needing_a_person
    if need:
        p("")
        p("Names a person has to settle — reported, never merged:")
        for r in need:
            w = r.who
            if w.how == "candidate":
                p(f"  {w.zoom:<28} ?  {w.candidates[0]}")
            elif w.how == "ambiguous":
                p(f"  {w.zoom:<28} ?  {' OR '.join(w.candidates)}   (matched to neither)")
            else:
                p(f"  {w.zoom:<28}    no roster name resembles it")
    p("")
    p("Details (Elaine's column, call-log side only):")
    for r in report.rows:
        for line in r.details:
            p(f"  {r.name:<28}{line}")
    return "\n".join(out)


def to_json(report: Report) -> dict:
    """The report as data, for the step that merges the declared numbers in."""
    def one(r: Row) -> dict:
        s = r.shift
        return {
            "zoom_name": r.who.zoom, "roster_name": r.who.roster, "match": r.who.how,
            "candidates": list(r.who.candidates),
            "calls": s.attempts, "answered": s.answered, "completes": s.completes(),
            "away_minutes": int(s.idle_time().total_seconds() // 60),
            "worked_minutes": int(s.worked().total_seconds() // 60),
            "short_minutes": int(s.shortfall().total_seconds() // 60),
            "started": s.started_at().isoformat() if s.started_at() else None,
            "finished": s.finished_at().isoformat() if s.finished_at() else None,
            "details": r.details,
            # Filled in from #results-pacificlinkglobal, never from here.
            "declared_completed": None, "declared_calls": None, "discrepancy": None,
        }
    return {
        "day": report.day.isoformat(),
        "parameters": {"complete_seconds": COMPLETE_SECONDS, "idle_seconds": IDLE_SECONDS,
                       "shift_hours": SHIFT_HOURS},
        "rows": [one(r) for r in report.rows],
        "off_shift": [one(r) for r in report.off_shift],
        "rang_only": report.rang_only,
    }


ELAINE_COLUMNS = (
    "Yellow - discrepancy on CS  Orange - for confirmation",
    "Numbers of Completed Surveys in WhatsApp",
    "Numbers of Completed Surveys in Call Logs",
    "Total Number of Calls in WhatsApp",
    "Total Number of Calls in Call Logs",
    "Discrepancy Identified? Y/N",
    "Details",
)


def write_xlsx(report: Report, path: str) -> bool:
    """Elaine's seven columns. Returns False, writing nothing, without openpyxl."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    declared = PatternFill("solid", fgColor="FFF2CC")   # the half that is not ours to fill
    ws.append(list(ELAINE_COLUMNS))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.fill = PatternFill("solid", fgColor="D9D9D9")
        c.border = border
    ws.append([report.day.strftime("%B %-d, %Y")] + [""] * 6)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F2F2F2")
        c.border = border
    for r in report.rows:
        s = r.shift
        ws.append([r.name, None, s.completes(), None, s.attempts, None, "\n".join(r.details) or "-"])
        n = ws.max_row
        for col in range(1, 8):
            ws.cell(n, col).border = border
            ws.cell(n, col).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (2, 4, 6):
            ws.cell(n, col).fill = declared
    ws.append(["TOTAL", None, sum(r.shift.completes() for r in report.rows), None,
               sum(r.shift.attempts for r in report.rows), None, ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.border = border
    for col, width in zip("ABCDEFG", (26, 14, 14, 14, 14, 13, 72)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A3"
    wb.save(path)
    return True


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("day", nargs="?", help="YYYY-MM-DD, Manila. Default: today in Manila.")
    ap.add_argument("--xlsx", help="write Elaine's layout here")
    ap.add_argument("--json", help="write the report as data here")
    args = ap.parse_args(argv)
    day = date.fromisoformat(args.day) if args.day else today_manila()
    report = audit(day)
    print(render_text(report))
    if args.json:
        with open(args.json, "w") as fh:
            json.dump(to_json(report), fh, indent=2)
        print(f"\njson: {args.json}")
    if args.xlsx:
        if write_xlsx(report, args.xlsx):
            print(f"xlsx: {args.xlsx}")
        else:
            print("xlsx: skipped — openpyxl is not installed", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
