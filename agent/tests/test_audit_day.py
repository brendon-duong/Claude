import json
import os
import tempfile
import unittest
from datetime import date

from business_agent.audit_day import audit, render_text, to_json, write_xlsx

DAY = date(2026, 9, 13)


def row(owner, when, seconds=0, result="Call Cancel", direction="outbound", caller=None):
    return {
        "caller_name": caller if caller is not None else (owner if direction == "outbound" else "Anonymous"),
        "caller_number": "+63288001", "callee_number": "+6421234567",
        "date_time": when, "duration": seconds, "result": result, "direction": direction,
        "owner": {"name": owner},
    }


def stub(rows):
    return lambda day, token: rows


class TestOneDay(unittest.TestCase):
    def setUp(self):
        # 05:30 UTC is 1:30pm Manila. Cherry works a full afternoon; Princess
        # makes one test call at 10:37 Manila and never comes back.
        rows = [
            row("Cherry Jean Raagas", f"2026-09-13T0{5 + (30 + m) // 60}:{(30 + m) % 60:02d}:00Z",
                seconds=160, result="Auto Recorded")
            for m in range(0, 181, 3)
        ]
        rows.append(row("Cherry Jean Raagas", "2026-09-13T07:00:30Z", direction="inbound"))
        # A missed call to her desk at 7:53am Manila, and one at 6:18pm.
        rows.append(row("Cherry Jean Raagas", "2026-09-12T23:53:00Z", direction="inbound", result="No Answer"))
        rows.append(row("Cherry Jean Raagas", "2026-09-13T10:18:00Z", direction="inbound", result="No Answer"))
        # An extension that only ever rang, unanswered. Not a caller.
        rows.append(row("Yvonne Eusebio", "2026-09-13T01:00:00Z", direction="inbound", result="No Answer"))
        rows.append(row("Princess Matildo", "2026-09-13T02:37:00Z", seconds=3, result="Auto Recorded"))
        rows.append(row("Chary Jay Sanchez", "2026-09-13T02:37:01Z", seconds=8, result="Auto Recorded"))
        self.report = audit(DAY, fetch=stub(rows))

    def test_the_one_who_worked_is_scored(self):
        self.assertEqual([r.name for r in self.report.rows], ["Cherry Jean Raagas"])
        self.assertEqual(self.report.rows[0].shift.completes(), 61)

    def test_the_inbound_calls_count_toward_her_total_and_not_a_phantom(self):
        self.assertEqual(self.report.rows[0].shift.attempts, 64)
        self.assertNotIn("Anonymous", render_text(self.report))

    def test_missed_calls_to_her_desk_do_not_touch_her_time(self):
        shift = self.report.rows[0].shift
        self.assertEqual(shift.shortfall().total_seconds(), 0)
        self.assertNotIn("6:18", render_text(self.report))

    def test_an_extension_that_only_rang_is_counted_not_named(self):
        self.assertEqual(self.report.rang_only, 1)
        self.assertNotIn("Yvonne", render_text(self.report))
        self.assertIn("1 extension(s) only received calls nobody answered", render_text(self.report))

    def test_the_test_callers_are_listed_not_scored(self):
        self.assertEqual(sorted(r.who.zoom for r in self.report.off_shift),
                         ["Chary Jay Sanchez", "Princess Matildo"])
        text = render_text(self.report)
        self.assertIn("listed, not scored", text)
        self.assertNotIn("13:30-10:37", text)

    def test_the_unmatched_name_is_reported(self):
        self.assertIn("Chary Jay Sanchez", [r.who.zoom for r in self.report.needing_a_person])
        self.assertIn("no roster name resembles it", render_text(self.report))

    def test_json_leaves_the_declared_half_empty(self):
        data = to_json(self.report)
        self.assertEqual(data["day"], "2026-09-13")
        self.assertIsNone(data["rows"][0]["declared_completed"])
        self.assertIsNone(data["rows"][0]["discrepancy"])
        json.dumps(data)  # serialisable

    def test_xlsx_has_elaines_seven_columns(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "a.xlsx")
            self.assertTrue(write_xlsx(self.report, path))
            ws = openpyxl.load_workbook(path).active
            self.assertEqual(ws.max_column, 7)
            self.assertEqual(ws["A2"].value, "September 13, 2026")
            self.assertEqual(ws["C3"].value, 61)     # completes from the logs
            self.assertIsNone(ws["B3"].value)        # declared: not ours to fill
